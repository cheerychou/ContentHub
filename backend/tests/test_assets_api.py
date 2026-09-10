import base64
import json
import subprocess

import pytest
from fastapi.testclient import TestClient

import app.routers.assets as assets_mod
from app.db import get_db
from app.main import app
from app.storage import FakeStorage, get_storage

# 1×1 红点 PNG（真字节，同 M2 冒烟 / test_media_attrs）
PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1Pe"
    "AAAADElEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC"
)


@pytest.fixture()
def fake():
    return FakeStorage()


@pytest.fixture()
def client(db_session, fake):
    def _override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = _override_get_db
    app.dependency_overrides[get_storage] = lambda: fake
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _md_file(name="文章.md", text="# 标题\n\n168 元保养 299 元漆面。"):
    return {"file": (name, text.encode("utf-8"), "text/markdown")}


def test_upload_master_defaults_to_drafting(client):
    resp = client.post(
        "/api/assets",
        data={"zone": "master", "title": "跨品牌售后"},
        files=_md_file(),
    )
    assert resp.status_code == 201
    body = resp.json()
    assert body["zone"] == "master"
    assert body["status"] == "drafting"
    assert body["content_type"] == "markdown"


def test_upload_publish_zone_rejected(client):
    resp = client.post(
        "/api/assets",
        data={"zone": "publish", "title": "公众号版"},
        files=_md_file(),
    )
    assert resp.status_code == 422


def test_upload_source_defaults_to_available(client):
    resp = client.post(
        "/api/assets",
        data={"zone": "source", "title": "信源"},
        files=_md_file(),
    )
    assert resp.status_code == 201
    assert resp.json()["status"] == "available"


def test_upload_topic_defaults_to_candidate(client):
    """选题策划页直传：zone=topic 默认候选（评审修复回归）。"""
    resp = client.post(
        "/api/assets",
        data={"zone": "topic", "title": "选题：途虎供应链"},
        files=_md_file(),
    )
    assert resp.status_code == 201
    body = resp.json()
    assert body["zone"] == "topic"
    assert body["status"] == "candidate"


def test_get_topic_asset_detail_includes_file_url(client):
    """topic 资产详情可取预签名 URL（storage.ZONES 缺 topic 曾致 500，评审修复回归）。"""
    create = client.post(
        "/api/assets",
        data={"zone": "topic", "title": "选题：蓝鲸"},
        files=_md_file("t.md"),
    )
    asset_id = create.json()["id"]
    resp = client.get(f"/api/assets/{asset_id}")
    assert resp.status_code == 200
    body = resp.json()
    assert body["zone"] == "topic"
    assert body["file_url"] == f"fake://topic/{asset_id}/t.md"


def test_external_link_source(client):
    resp = client.post(
        "/api/assets/external",
        json={
            "zone": "source",
            "title": "九部委 17 条政策原文",
            "source_url": "https://example.com/policy",
        },
    )
    assert resp.status_code == 201
    assert resp.json()["content_type"] == "link"


def test_search_by_body_text(client):
    client.post("/api/assets", data={"zone": "master", "title": "甲文"},
                files=_md_file("a.md", "正文含 关键词蓝鲸。"))
    client.post("/api/assets", data={"zone": "master", "title": "乙文"},
                files=_md_file("b.md", "正文无关。"))
    resp = client.get("/api/assets", params={"q": "蓝鲸"})
    assert resp.status_code == 200
    titles = [a["title"] for a in resp.json()]
    assert titles == ["甲文"]


def test_status_transition_valid_and_invalid(client):
    create = client.post("/api/assets", data={"zone": "master", "title": "t"},
                         files=_md_file())
    asset_id = create.json()["id"]

    bad = client.patch(f"/api/assets/{asset_id}/status", json={"status": "published"})
    assert bad.status_code == 422

    ok = client.patch(f"/api/assets/{asset_id}/status", json={"status": "finalized"})
    assert ok.status_code == 200
    assert ok.json()["status"] == "finalized"


def test_delete_asset(client):
    create = client.post("/api/assets", data={"zone": "master", "title": "t"},
                         files=_md_file())
    asset_id = create.json()["id"]
    assert client.delete(f"/api/assets/{asset_id}").status_code == 204
    assert client.get(f"/api/assets/{asset_id}").status_code == 404


def test_upload_docx_extracts_text(client):
    from io import BytesIO
    from docx import Document
    buf = BytesIO()
    doc = Document()
    doc.add_paragraph("第一段：途虎学什么。")
    doc.add_paragraph("第二段：供应链视角。")
    doc.save(buf)
    resp = client.post(
        "/api/assets", data={"zone": "master", "title": "Word 稿"},
        files={"file": ("途虎.docx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 201
    body = resp.json()
    assert body["content_type"] == "docx"
    assert "第二段：供应链视角。" in body["text_content"]


def test_upload_broken_docx_422_and_no_object(client, fake):
    """损坏 docx → 422，且先抽取后落盘：存储中不残留孤儿对象。"""
    resp = client.post(
        "/api/assets", data={"zone": "master", "title": "坏 Word"},
        files={"file": ("x.docx", b"not a docx",
                        "application/vnd.openxmlformats-officedocument"
                        ".wordprocessingml.document")},
    )
    assert resp.status_code == 422
    assert "docx 解析失败" in resp.json()["detail"]
    assert fake.objects == {}


# ---------- M7 素材属性抽取接入 ----------

def test_upload_image_extracts_attrs(client):
    """上传图片 → 201，meta.attrs 含格式与宽高。"""
    resp = client.post(
        "/api/assets", data={"zone": "master", "title": "封面图"},
        files={"file": ("cover.png", PNG_1X1, "image/png")},
    )
    assert resp.status_code == 201
    assert resp.json()["meta"]["attrs"] == {
        "format": "PNG", "width": 1, "height": 1}


def test_upload_md_extracts_word_count(client):
    """上传 markdown → meta.attrs 字数与语言。"""
    resp = client.post(
        "/api/assets", data={"zone": "master", "title": "文稿"},
        files=_md_file("文稿.md", "# 标题\n\n途虎养车供应链视角。"),
    )
    assert resp.status_code == 201
    attrs = resp.json()["meta"]["attrs"]
    assert attrs["word_count"] > 0
    assert attrs["language"] == "zh"


def test_upload_broken_image_no_attrs_still_201(client):
    """失败隔离：垃圾图片字节 → 201 且 meta 无 attrs，上传不失败。"""
    resp = client.post(
        "/api/assets", data={"zone": "master", "title": "坏图"},
        files={"file": ("broken.png", b"definitely not png", "image/png")},
    )
    assert resp.status_code == 201
    assert "attrs" not in resp.json()["meta"]


def test_upload_large_video_spool_path_extracts_attrs(client, monkeypatch):
    """大文件（spool）路径：视频另落临时文件喂 ffprobe（canned JSON），写入 attrs。"""
    monkeypatch.setattr(assets_mod, "INLINE_TEXT_LIMIT", 8)  # 强制走 spool 大文件路径

    def fake_run(cmd, **kwargs):
        return subprocess.CompletedProcess(
            cmd, 0, stdout=json.dumps({
                "streams": [{"codec_type": "video", "codec_name": "h264",
                             "width": 1920, "height": 1080}],
                "format": {"duration": "12.5"},
            }), stderr="")

    monkeypatch.setattr(assets_mod.media_attrs.subprocess, "run", fake_run)
    resp = client.post(
        "/api/assets", data={"zone": "master", "title": "大视频"},
        files={"file": ("demo.mp4", b"\x00" * 16, "video/mp4")},
    )
    assert resp.status_code == 201
    assert resp.json()["meta"]["attrs"] == {
        "format": "h264", "width": 1920, "height": 1080, "duration_seconds": 12.5}


def test_upload_small_video_memory_path_extracts_attrs(client, monkeypatch):
    """小视频（<50MB 常见）走内存路径也要抽 attrs（真实冒烟回归：曾漏分派）。"""

    def fake_run(cmd, **kwargs):
        return subprocess.CompletedProcess(
            cmd, 0, stdout=json.dumps({
                "streams": [{"codec_type": "video", "codec_name": "h264",
                             "width": 320, "height": 240}],
                "format": {"duration": "1.0"},
            }), stderr="")

    monkeypatch.setattr(assets_mod.media_attrs.subprocess, "run", fake_run)
    resp = client.post(
        "/api/assets", data={"zone": "master", "title": "小视频"},
        files={"file": ("demo.mp4", b"\x00" * 16, "video/mp4")},
    )
    assert resp.status_code == 201
    assert resp.json()["meta"]["attrs"] == {
        "format": "h264", "width": 320, "height": 240, "duration_seconds": 1.0}


# ---------- M7 人工补录 attrs PATCH ----------

def _upload_md(client, text="# 标题\n\n途虎养车供应链视角，共两句。"):
    resp = client.post(
        "/api/assets", data={"zone": "master", "title": "文稿"},
        files=_md_file("文稿.md", text),
    )
    assert resp.status_code == 201
    return resp.json()


def test_patch_attrs_adds_manual_key_preserving_auto_keys(client):
    asset = _upload_md(client)
    auto = asset["meta"]["attrs"]
    resp = client.patch(f"/api/assets/{asset['id']}/attrs",
                        json={"attrs": {"author": "周大波"}})
    assert resp.status_code == 200
    attrs = resp.json()["meta"]["attrs"]
    assert attrs["author"] == "周大波"          # 补录键写入
    assert attrs["word_count"] == auto["word_count"]  # 自动键保留
    assert attrs["language"] == auto["language"]


def test_patch_attrs_manual_key_wins_on_conflict(client):
    """补录键与自动键同名时以补录为准（用户人工输入优先）。"""
    asset = _upload_md(client)
    resp = client.patch(f"/api/assets/{asset['id']}/attrs",
                        json={"attrs": {"language": "en"}})
    assert resp.status_code == 200
    attrs = resp.json()["meta"]["attrs"]
    assert attrs["language"] == "en"
    assert attrs["word_count"] == asset["meta"]["attrs"]["word_count"]


def test_patch_attrs_deep_merges_nested_dicts(client):
    asset = _upload_md(client)
    asset_id = asset["id"]
    assert client.patch(f"/api/assets/{asset_id}/attrs",
                        json={"attrs": {"manual": {"platform": "wx"}}}
                        ).status_code == 200
    resp = client.patch(f"/api/assets/{asset_id}/attrs",
                        json={"attrs": {"manual": {"author": "周"}}})
    assert resp.status_code == 200
    assert resp.json()["meta"]["attrs"]["manual"] == {
        "platform": "wx", "author": "周"}  # 深合并：嵌套键互不覆盖


def test_patch_attrs_empty_object_rejected_422(client):
    asset = _upload_md(client)
    for bad in ({"attrs": {}}, {"attrs": "不是对象"}, {}):
        resp = client.patch(f"/api/assets/{asset['id']}/attrs", json=bad)
        assert resp.status_code == 422, bad


def test_patch_attrs_404_for_missing_asset(client):
    resp = client.patch(
        "/api/assets/00000000-0000-0000-0000-000000000000/attrs",
        json={"attrs": {"author": "周大波"}},
    )
    assert resp.status_code == 404


# ---------- M7 列表按素材类型过滤（content_type 逗号分隔多值） ----------

def _upload_docx(client, title="Word 稿", name="稿.docx"):
    from io import BytesIO
    from docx import Document
    buf = BytesIO()
    doc = Document()
    doc.add_paragraph("第一段：途虎学什么。")
    doc.save(buf)
    resp = client.post(
        "/api/assets", data={"zone": "master", "title": title},
        files={"file": (name, buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument"
                        ".wordprocessingml.document")},
    )
    assert resp.status_code == 201
    return resp.json()


def test_list_filter_content_type(client):
    """M7 类型筛选：md+docx+png 三资产——多值命中两类、单值命中一类、缺省返回全部。"""
    client.post("/api/assets", data={"zone": "master", "title": "md 文稿"},
                files=_md_file("a.md"))
    _upload_docx(client, title="docx 文稿", name="b.docx")
    client.post("/api/assets", data={"zone": "master", "title": "png 封面"},
                files={"file": ("c.png", PNG_1X1, "image/png")})

    def types(params):
        resp = client.get("/api/assets", params=params)
        assert resp.status_code == 200
        return sorted(a["content_type"] for a in resp.json())

    # 多值（逗号分隔）：恰命中 markdown 与 docx 两类，不含 png
    assert types({"zone": "master", "content_type": "markdown,docx"}) \
        == ["docx", "markdown"]
    # 单值：仅命中一类
    assert types({"zone": "master", "content_type": "image"}) == ["image"]
    # 缺省：不过滤，三类齐全
    assert types({"zone": "master"}) == ["docx", "image", "markdown"]
